"""The authoritative OOXML package parser.

This module deliberately does not use ``python-pptx`` to discover package
parts. The ZIP and relationship layers are parsed directly so unsupported
PresentationML is still retained and visible in the report.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict
import hashlib
import importlib.metadata
from io import BytesIO
import math
import posixpath
from pathlib import Path, PurePosixPath, PureWindowsPath
import re
from typing import Any, Iterable
from urllib.parse import unquote, urlsplit
import zipfile

from defusedxml import ElementTree as SafeET

from .geometry import IDENTITY, child_transform_matrix, geometry_for_object
from .diagrams import add_native_diagram_evidence
from .models import (
    DECKIR_SCHEMA,
    DECKIR_SCHEMA_VERSION,
    DeckIR,
    ExtractionReport,
    MediaRecord,
    PartRecord,
    RelationshipRecord,
    SlideRecord,
)
from .semantics import native_semantics, resolve_style
from .visual import add_native_visual_evidence

VERSION = "0.5.0"
CONTENT_TYPES = "[Content_Types].xml"
RELATIONSHIP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
STRICT_RELATIONSHIP_NS = "http://purl.oclc.org/ooxml/officeDocument/relationships"
STRICT_OFFICE_DOCUMENT_TYPE = "http://purl.oclc.org/ooxml/officeDocument"
STRICT_RELATIONSHIP_TYPE_NS = "http://purl.oclc.org/ooxml"
PACKAGE_RELATIONSHIP_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
STRICT_PACKAGE_RELATIONSHIP_NS = "http://purl.oclc.org/ooxml/package/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
STRICT_CONTENT_TYPES_NS = "http://purl.oclc.org/ooxml/package/content-types"
PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
STRICT_PRESENTATION_NS = "http://purl.oclc.org/ooxml/presentationml/main"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
STRICT_DRAWING_NS = "http://purl.oclc.org/ooxml/drawingml/main"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
STRICT_CHART_NS = "http://purl.oclc.org/ooxml/drawingml/chart"
DIAGRAM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
STRICT_DIAGRAM_NS = "http://purl.oclc.org/ooxml/drawingml/diagram"
MARKUP_COMPATIBILITY_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
STRICT_MARKUP_COMPATIBILITY_NS = "http://purl.oclc.org/ooxml/markup-compatibility/2006"
RELATIONSHIP_NAMESPACES = frozenset({RELATIONSHIP_NS, STRICT_RELATIONSHIP_NS})
PACKAGE_RELATIONSHIP_NAMESPACES = frozenset({PACKAGE_RELATIONSHIP_NS, STRICT_PACKAGE_RELATIONSHIP_NS})
CONTENT_TYPES_NAMESPACES = frozenset({CONTENT_TYPES_NS, STRICT_CONTENT_TYPES_NS})
MARKUP_COMPATIBILITY_NAMESPACES = frozenset({MARKUP_COMPATIBILITY_NS, STRICT_MARKUP_COMPATIBILITY_NS})
OOXML_NAMESPACES = frozenset(
    {
        PRESENTATION_NS,
        STRICT_PRESENTATION_NS,
        DRAWING_NS,
        STRICT_DRAWING_NS,
        CHART_NS,
        STRICT_CHART_NS,
        DIAGRAM_NS,
        STRICT_DIAGRAM_NS,
        *RELATIONSHIP_NAMESPACES,
        *PACKAGE_RELATIONSHIP_NAMESPACES,
        *CONTENT_TYPES_NAMESPACES,
        *MARKUP_COMPATIBILITY_NAMESPACES,
    }
)
_SLIDE_NUMBER = re.compile(r"^ppt/slides/slide(\d+)\.xml$")
_SUPPORTED_MC_REQUIREMENTS = frozenset({"a", "c", "dgm", "p", "r"})


class ExtractionError(ValueError):
    """Raised when the input is not a safe, readable source document."""


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _namespace(tag: str) -> str:
    return tag[1 : tag.index("}")] if tag.startswith("{") and "}" in tag else ""


def _known_namespace(tag: str, allowed: frozenset[str] = OOXML_NAMESPACES) -> bool:
    namespace = _namespace(tag)
    return not namespace or namespace in allowed


def _attr(element: Any, local_name: str) -> str | None:
    direct = element.attrib.get(local_name)
    if direct is not None:
        return direct
    allowed = RELATIONSHIP_NAMESPACES if local_name == "id" else frozenset()
    for key, value in element.attrib.items():
        if _local_name(key) == local_name and _namespace(key) in allowed:
            return value
    return None


def _relationship_attr(element: Any, local_name: str) -> str | None:
    for key, value in element.attrib.items():
        if _local_name(key) == local_name and _namespace(key) in RELATIONSHIP_NAMESPACES:
            return value
    return None


def _parse_xml(data: bytes, part: str) -> Any:
    try:
        return SafeET.fromstring(data)
    except Exception as exc:
        raise ExtractionError(f"Invalid XML in package part {part}: {exc}") from exc


def _parse_optional_xml(data: bytes | None, part: str, warnings: list[str]) -> Any | None:
    if data is None:
        return None
    try:
        return _parse_xml(data, part)
    except ExtractionError as exc:
        warnings.append(str(exc))
        return None


def _safe_part_name(name: str, *, directory: bool = False) -> str:
    if not isinstance(name, str) or not name or name.startswith("/") or "\\" in name:
        raise ExtractionError(f"Unsafe package part name: {name!r}")
    if any(ord(character) < 0x20 or ord(character) == 0x7F for character in name):
        raise ExtractionError(f"Unsafe package part name: {name!r}")
    if PureWindowsPath(name).drive or re.match(r"^[A-Za-z]:", name) or name.startswith("//"):
        raise ExtractionError(f"Unsafe package part name: {name!r}")
    parts = name.split("/")
    if directory and parts[-1] == "":
        parts.pop()
    if not parts or any(piece in ("", ".", "..") for piece in parts):
        raise ExtractionError(f"Unsafe package part name: {name!r}")
    return name


def _relationship_source(name: str) -> str:
    path = PurePosixPath(name)
    if path.name == ".rels" and path.parent.name == "_rels":
        owner = path.parent.parent
        return "" if str(owner) == "." else str(owner)
    if path.parent.name != "_rels":
        raise ExtractionError(f"Invalid relationships part name: {name}")
    return str(path.parent.parent / path.stem)


def _resolve_target(source: str, target: str, mode: str | None) -> str | None:
    normalized_mode = mode.strip().lower() if mode is not None else None
    if normalized_mode == "external" or not target:
        return None
    if normalized_mode is not None and normalized_mode != "internal":
        return None
    if not isinstance(target, str) or "\\" in target or any(ord(character) < 0x20 for character in target):
        return None
    try:
        parsed = urlsplit(target)
    except ValueError:
        return None
    if parsed.scheme or parsed.netloc:
        return None
    path = unquote(parsed.path)
    if not path:
        return None
    if path.startswith("/"):
        if ".." in path.split("/"):
            return None
        resolved = posixpath.normpath(path.lstrip("/"))
    else:
        base = PurePosixPath(source).parent if source else PurePosixPath(".")
        resolved = posixpath.normpath(posixpath.join(str(base), path))
    if resolved in {"", ".", ".."} or any(piece == ".." for piece in resolved.split("/")):
        return None
    try:
        return _safe_part_name(resolved)
    except ExtractionError:
        return None


def _relationship_type_name(value: str | None) -> str:
    raw = (value or "").strip().rstrip("/")
    if raw.startswith("/"):
        raw = raw[1:]
    if raw in {
        "chart",
        "diagram",
        "diagramData",
        "hyperlink",
        "image",
        "notesSlide",
        "officeDocument",
        "oleObject",
        "slide",
        "slideLayout",
        "slideMaster",
        "theme",
    }:
        return raw.casefold()
    if value and value.strip().rstrip("/") == STRICT_OFFICE_DOCUMENT_TYPE:
        return "officedocument"
    for namespace in (*RELATIONSHIP_NAMESPACES, STRICT_RELATIONSHIP_TYPE_NS):
        prefix = f"{namespace}/"
        if value and value.strip().rstrip("/").startswith(prefix):
            candidate = value.strip().rstrip("/")[len(prefix) :]
            if candidate in {
                "chart",
                "diagram",
                "diagramData",
                "hyperlink",
                "image",
                "notesSlide",
                "officeDocument",
                "oleObject",
                "slide",
                "slideLayout",
                "slideMaster",
                "theme",
            }:
                return candidate.casefold()
    return ""


def _relationship_matches(relation: RelationshipRecord, suffix: str) -> bool:
    return _relationship_type_name(relation.relationship_type) == _relationship_type_name(suffix)


def _relationship_map(relations: Iterable[RelationshipRecord]) -> dict[str, RelationshipRecord]:
    result: dict[str, RelationshipRecord] = {}
    for relation in relations:
        result.setdefault(relation.relationship_id, relation)
    return result


def _dependency_versions() -> dict[str, str]:
    versions: dict[str, str] = {}
    for distribution in ("defusedxml", "lxml", "python-pptx"):
        try:
            versions[distribution] = importlib.metadata.version(distribution)
        except importlib.metadata.PackageNotFoundError:
            versions[distribution] = "not-installed"
    return versions


def _content_types(raw: bytes | None, warnings: list[str] | None = None) -> dict[str, str]:
    if raw is None:
        return {}
    root = _parse_xml(raw, CONTENT_TYPES)
    if _local_name(root.tag) != "Types" or not _known_namespace(root.tag, CONTENT_TYPES_NAMESPACES):
        raise ExtractionError(f"Invalid content types part {CONTENT_TYPES}: expected Types root")
    result: dict[str, str] = {}
    for item in root:
        name = _local_name(item.tag)
        if not _known_namespace(item.tag, CONTENT_TYPES_NAMESPACES):
            continue
        if name == "Override":
            part_name, content_type = item.get("PartName"), item.get("ContentType")
            if not part_name or not content_type:
                if warnings is not None:
                    warnings.append("Content type Override is missing PartName or ContentType")
                continue
            if not part_name.startswith("/"):
                if warnings is not None:
                    warnings.append(f"Content type Override has an invalid PartName: {part_name!r}")
                continue
            normalized = unquote(part_name[1:])
            try:
                normalized = _safe_part_name(normalized)
            except ExtractionError:
                if warnings is not None:
                    warnings.append(f"Content type Override has an unsafe PartName: {part_name!r}")
                continue
            if normalized in result and warnings is not None:
                warnings.append(f"Duplicate content type Override: {normalized}")
            if normalized in result:
                continue
            result[normalized] = content_type
        elif name == "Default":
            extension, content_type = item.get("Extension"), item.get("ContentType")
            if extension and content_type:
                key = f"*.{extension.lower()}"
                if key in result and warnings is not None:
                    warnings.append(f"Duplicate content type Default: {extension}")
                if key not in result:
                    result[key] = content_type
            elif warnings is not None:
                warnings.append("Content type Default is missing Extension or ContentType")
    return result


def _part_content_type(name: str, content_types: dict[str, str]) -> str:
    if name in content_types:
        return content_types[name]
    suffix = Path(name).suffix.lower().lstrip(".")
    return content_types.get(f"*.{suffix}", "application/octet-stream")


def _relationships(
    names_to_bytes: dict[str, bytes], warnings: list[str]
) -> tuple[list[RelationshipRecord], dict[str, list[RelationshipRecord]]]:
    records: list[RelationshipRecord] = []
    by_source: dict[str, list[RelationshipRecord]] = defaultdict(list)
    for name in sorted(names_to_bytes):
        if not name.endswith(".rels") or "/_rels/" not in f"/{name}":
            continue
        try:
            source = _relationship_source(name)
            root = _parse_xml(names_to_bytes[name], name)
            if _local_name(root.tag) != "Relationships" or not _known_namespace(root.tag, PACKAGE_RELATIONSHIP_NAMESPACES):
                raise ExtractionError(f"Invalid relationships part {name}: expected Relationships root")
        except ExtractionError as exc:
            warnings.append(str(exc))
            continue
        seen_ids: set[str] = set()
        for relation in root:
            if _local_name(relation.tag) != "Relationship" or not _known_namespace(relation.tag, PACKAGE_RELATIONSHIP_NAMESPACES):
                continue
            relationship_id = relation.get("Id", "")
            relationship_type = relation.get("Type", "")
            target = relation.get("Target", "")
            if not relationship_id:
                warnings.append(f"Relationship in {name} has no Id")
                continue
            if relationship_id in seen_ids:
                warnings.append(f"Duplicate relationship Id in {name}: {relationship_id}")
            seen_ids.add(relationship_id)
            if not relationship_type:
                warnings.append(f"Relationship {relationship_id} in {name} has no Type")
            if not target:
                warnings.append(f"Relationship {relationship_id} in {name} has no Target")
            target_mode = relation.get("TargetMode")
            normalized_target_mode = target_mode.strip().lower() if target_mode is not None else None
            if normalized_target_mode is not None and normalized_target_mode not in {"internal", "external"}:
                warnings.append(f"Relationship {relationship_id} in {name} has an invalid TargetMode: {target_mode!r}")
            item = RelationshipRecord(
                source_part=source,
                relationship_id=relationship_id,
                relationship_type=relationship_type,
                target=target,
                target_mode=target_mode,
                resolved_target=_resolve_target(source, target, target_mode),
            )
            records.append(item)
            by_source[source].append(item)
            if target and normalized_target_mode not in {"external"}:
                if item.resolved_target is None:
                    warnings.append(f"Relationship {relationship_id} in {name} has an unsafe or unresolvable target: {target!r}")
                elif item.resolved_target not in names_to_bytes:
                    warnings.append(f"Relationship {relationship_id} in {name} points to missing part: {item.resolved_target}")
    records.sort(key=lambda item: (item.source_part, item.relationship_id, item.relationship_type, item.target))
    return records, by_source


def _rels_to(by_source: dict[str, list[RelationshipRecord]], source: str, suffix: str) -> str | None:
    for relation in by_source.get(source, []):
        if _relationship_matches(relation, suffix):
            return relation.resolved_target
    return None


def _texts(root: Any) -> list[str]:
    return [
        element.text
        for element in root.iter()
        if _local_name(element.tag) == "t"
        and _known_namespace(element.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS, DRAWING_NS, STRICT_DRAWING_NS}))
        and element.text
    ]


def _hyperlinks(root: Any, relations: Iterable[RelationshipRecord]) -> list[dict[str, Any]]:
    relation_map = _relationship_map(relations)
    found: list[dict[str, Any]] = []
    for element in root.iter():
        kind = _local_name(element.tag)
        if kind not in {"hlinkClick", "hlinkHover"} or not _known_namespace(
            element.tag,
            frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS, DRAWING_NS, STRICT_DRAWING_NS}),
        ):
            continue
        relationship_id = _relationship_attr(element, "id")
        relation = relation_map.get(relationship_id or "")
        found.append(
            {
                "kind": kind,
                "relationship_id": relationship_id,
                "target": relation.target if relation else None,
                "resolved_target": relation.resolved_target if relation else None,
                "action": element.get("action"),
            }
        )
    return found


def _alt_text(root: Any) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for element in root.iter():
        if _local_name(element.tag) != "cNvPr" or not _known_namespace(
            element.tag,
            frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS, DRAWING_NS, STRICT_DRAWING_NS}),
        ):
            continue
        descr, title = element.get("descr"), element.get("title")
        if descr or title:
            result.append({"id": element.get("id", ""), "name": element.get("name", ""), "descr": descr or "", "title": title or ""})
    return result


def _animations(root: Any) -> list[dict[str, Any]]:
    animation_names = {"timing", "par", "seq", "anim", "set", "animEffect", "animMotion", "cmd"}
    result: list[dict[str, Any]] = []
    for element in root.iter():
        name = _local_name(element.tag)
        if name in animation_names and _known_namespace(element.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS})):
            result.append({"element": name, "attributes": dict(element.attrib)})
    return result


def _first_descendant(root: Any, local_name: str) -> Any | None:
    return next(
        (
            item
            for item in root.iter()
            if _local_name(item.tag) == local_name and _known_namespace(item.tag)
        ),
        None,
    )


def _slide_dimensions(root: Any | None) -> tuple[float, float]:
    if root is not None:
        size = _first_descendant(root, "sldSz")
        if size is not None:
            if size.get("cx") is None or size.get("cy") is None:
                raise ExtractionError("Invalid slide dimensions in presentation.xml")
            try:
                width, height = float(size.get("cx")), float(size.get("cy"))
            except (TypeError, ValueError) as exc:
                raise ExtractionError("Invalid slide dimensions in presentation.xml") from exc
            if not math.isfinite(width) or not math.isfinite(height) or width <= 0 or height <= 0:
                raise ExtractionError("Invalid slide dimensions in presentation.xml")
            return width, height
    # ISO/IEC 29500 default widescreen-independent slide size (10 x 7.5 in).
    return 9_144_000.0, 6_858_000.0


def _xml_path(parent_path: str, element: Any, siblings: list[Any]) -> str:
    local = _local_name(element.tag)
    same_name = [item for item in siblings if _local_name(item.tag) == local]
    return f"{parent_path}/{local}[{same_name.index(element) + 1}]"


def _shape_id_element(element: Any) -> Any | None:
    return _first_descendant(element, "cNvPr")


def _placeholder_geometry(element: Any, root: Any | None) -> Any | None:
    placeholder = _first_descendant(element, "ph")
    if placeholder is None or root is None:
        return None
    wanted_type = placeholder.get("type", "body")
    wanted_idx = placeholder.get("idx")
    for candidate in root.iter():
        if _local_name(candidate.tag) not in {"sp", "pic", "graphicFrame"} or not _known_namespace(candidate.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS})):
            continue
        candidate_placeholder = _first_descendant(candidate, "ph")
        if candidate_placeholder is None or candidate_placeholder.get("type", "body") != wanted_type:
            continue
        if wanted_idx is not None and candidate_placeholder.get("idx") != wanted_idx:
            continue
        if _first_descendant(candidate, "xfrm") is not None:
            return candidate
    return None


def _object_type(element: Any) -> str | None:
    if not _known_namespace(element.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS})):
        return None
    local = _local_name(element.tag)
    if local == "sp":
        return "text" if _first_descendant(element, "txBody") is not None else "shape"
    if local == "pic":
        return "image"
    if local == "cxnSp":
        return "connector"
    if local == "grpSp":
        return "group"
    if local == "graphicFrame":
        graphic_data = _first_descendant(element, "graphicData")
        uri = graphic_data.get("uri", "") if graphic_data is not None else ""
        descendants = {_local_name(item.tag) for item in element.iter() if _known_namespace(item.tag)}
        if "table" in descendants or uri.endswith("/table"):
            return "table"
        if "chart" in descendants or "chart" in uri:
            return "chart"
        if "relIds" in descendants or "diagram" in uri:
            return "smartart"
        return "shape"
    if local in {"contentPart", "oleObj"}:
        return "shape"
    return None


def _native_shape_type(element: Any, semantic_type: str | None) -> str | None:
    """Preserve the OOXML/Python-pptx shape subtype beside the semantic type."""
    if not _known_namespace(element.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS})):
        return None
    local = _local_name(element.tag)
    if local == "sp":
        if _first_descendant(element, "ph") is not None:
            return "PLACEHOLDER"
        if _first_descendant(element, "custGeom") is not None:
            return "FREEFORM"
        return "TEXT_BOX" if _first_descendant(element, "txBody") is not None else "AUTO_SHAPE"
    if local == "pic":
        return "PICTURE"
    if local == "cxnSp":
        return "CONNECTOR"
    if local == "grpSp":
        return "GROUP"
    return {
        "table": "TABLE",
        "chart": "CHART",
        "smartart": "SMARTART",
    }.get(semantic_type, "GRAPHIC_FRAME" if local == "graphicFrame" else semantic_type)


def _object_relationship_ids(element: Any) -> list[str]:
    result: list[str] = []
    for descendant in element.iter():
        for key, value in descendant.attrib.items():
            if _namespace(key) in RELATIONSHIP_NAMESPACES and _local_name(key) in {"id", "embed", "link", "dm", "lo", "qs"}:
                if value not in result:
                    result.append(value)
    return result


def _object_style(element: Any) -> dict[str, Any]:
    placeholder = _first_descendant(element, "ph")
    style: dict[str, Any] = {}
    if placeholder is not None and placeholder.get("type"):
        style["placeholder_type"] = placeholder.get("type")
    if placeholder is not None and placeholder.get("idx"):
        style["placeholder_idx"] = placeholder.get("idx")
    transform = _first_descendant(element, "xfrm")
    if transform is not None and transform.get("rot"):
        try:
            rotation = float(transform.get("rot")) / 60_000.0
            if math.isfinite(rotation):
                style["rotation_degrees"] = rotation
        except (TypeError, ValueError):
            pass
    return style


def _object_text(element: Any) -> str:
    paragraphs: list[str] = []
    for paragraph in (item for item in element.iter() if _local_name(item.tag) == "p" and _known_namespace(item.tag)):
        pieces: list[str] = []
        for descendant in paragraph.iter():
            if not _known_namespace(descendant.tag):
                continue
            kind = _local_name(descendant.tag)
            if kind == "t" and descendant.text is not None:
                pieces.append(descendant.text)
            elif kind == "br":
                pieces.append("\n")
            elif kind == "tab":
                pieces.append("\t")
        paragraphs.append("".join(pieces))
    if paragraphs:
        return "\n".join(paragraphs)
    return "".join(_texts(element))


def _native_source(part: str | None, xml_path: str | None, confidence: float = 1.0) -> dict[str, Any]:
    return {
        "layer": "native_ooxml",
        "xml_part": part,
        "xml_path": xml_path,
        "confidence": confidence,
    }


def _slide_objects(
    root: Any,
    slide_id: str,
    part: str,
    slide_relations: Iterable[RelationshipRecord],
    slide_width: float,
    slide_height: float,
    asset_ids: dict[str, str],
    package_parts: dict[str, bytes],
    layout_root: Any | None,
    master_root: Any | None,
    theme_root: Any | None,
) -> list[dict[str, Any]]:
    shape_tree = _first_descendant(root, "spTree")
    if shape_tree is None:
        return []
    relation_map = _relationship_map(slide_relations)
    relation_ids = {
        relationship_id: f"{part}:{relationship_id}" for relationship_id in relation_map
    }
    objects: list[dict[str, Any]] = []
    ordinal = 0

    def visit(
        container: Any,
        container_path: str,
        parent_id: str | None,
        parent_matrix: tuple[float, float, float, float, float, float] = IDENTITY,
        transform_chain: list[str] | None = None,
    ) -> None:
        nonlocal ordinal
        ancestor_chain = transform_chain or []
        children = list(container)
        for child in children:
            kind = _object_type(child)
            child_path = _xml_path(container_path, child, children)
            if kind is None:
                local = _local_name(child.tag)
                if local == "AlternateContent" and _namespace(child.tag) in MARKUP_COMPATIBILITY_NAMESPACES:
                    alternatives = [
                        item
                        for item in list(child)
                        if _local_name(item.tag) in {"Choice", "Fallback"}
                        and _namespace(item.tag) in MARKUP_COMPATIBILITY_NAMESPACES
                    ]
                    selected = next(
                        (
                            item
                            for item in alternatives
                            if _local_name(item.tag) == "Choice"
                            and all(
                                requirement in _SUPPORTED_MC_REQUIREMENTS
                                for requirement in item.get("Requires", "").split()
                            )
                        ),
                        None,
                    )
                    if selected is None:
                        selected = next((item for item in alternatives if _local_name(item.tag) == "Fallback"), None)
                    if selected is not None:
                        visit(selected, child_path, parent_id, parent_matrix, ancestor_chain)
                elif local in {"Choice", "Fallback", "lockedCanvas"} and (
                    _namespace(child.tag) in MARKUP_COMPATIBILITY_NAMESPACES
                    or _known_namespace(child.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS}))
                ):
                    visit(child, child_path, parent_id, parent_matrix, ancestor_chain)
                continue
            ordinal += 1
            object_id = f"{slide_id}-shape-{ordinal:02d}"
            placeholder_fallback = _placeholder_geometry(child, layout_root)
            if placeholder_fallback is None:
                placeholder_fallback = _placeholder_geometry(child, master_root)
            bbox, geometry = geometry_for_object(
                child,
                slide_width,
                slide_height,
                parent_matrix,
                [*ancestor_chain, object_id],
                placeholder_fallback,
            )
            shape_id = _shape_id_element(child)
            object_relationships = [
                relation_ids.get(value, f"{part}:{value}")
                for value in _object_relationship_ids(child)
            ]
            object_record: dict[str, Any] = {
                "id": object_id,
                "slide_id": slide_id,
                "parent_id": parent_id,
                "type": kind,
                "shape_type": _native_shape_type(child, kind),
                "bbox": bbox,
                "z_order": ordinal,
                "text": _object_text(child),
                "style": _object_style(child),
                "geometry": geometry,
                "relationships": object_relationships,
                "source": _native_source(part, child_path),
            }
            raw_style, resolved_style, inherited_from = resolve_style(
                child,
                layout_root,
                master_root,
                theme_root,
            )
            object_record["raw_style"] = raw_style
            object_record["resolved_style"] = resolved_style
            object_record["inherited_from"] = inherited_from
            object_record.update(
                native_semantics(
                    child,
                    kind,
                    slide_relations,
                    package_parts,
                    asset_ids,
                )
            )
            if shape_id is not None:
                object_record["native_id"] = shape_id.get("id")
                object_record["name"] = shape_id.get("name", "")
            for relationship_id in _object_relationship_ids(child):
                relation = relation_map.get(relationship_id)
                if relation and relation.resolved_target in asset_ids:
                    object_record["asset_id"] = asset_ids[relation.resolved_target]
                    if relation.resolved_target.startswith("ppt/embeddings/"):
                        object_record["embedded"] = True
                    break
            objects.append(object_record)
            if kind == "group":
                visit(
                    child,
                    child_path,
                    object_id,
                    child_transform_matrix(child, parent_matrix),
                    [*ancestor_chain, object_id],
                )

    visit(shape_tree, "/sld[1]", None)
    return objects


def _python_pptx_summary(source: Path) -> dict[str, Any]:
    try:
        import pptx  # type: ignore
    except ImportError:
        return {"available": False, "reason": "python-pptx is not installed"}
    try:
        presentation = pptx.Presentation(str(source))
        slides: list[dict[str, Any]] = []
        for number, slide in enumerate(presentation.slides, 1):
            shapes: list[dict[str, Any]] = []
            for shape in slide.shapes:
                item: dict[str, Any] = {"name": shape.name, "shape_type": str(shape.shape_type)}
                if hasattr(shape, "text") and shape.text:
                    item["text"] = shape.text
                shapes.append(item)
            slides.append({"number": number, "shapes": shapes})
        return {"available": True, "version": getattr(pptx, "__version__", "unknown"), "slides": slides}
    except Exception as exc:  # enrichment must not hide package evidence
        return {"available": True, "error": str(exc)}


def _comments(names_to_bytes: dict[str, bytes], warnings: list[str] | None = None) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for name, raw in sorted(names_to_bytes.items()):
        if "/comments/" not in f"/{name}" or not name.endswith(".xml"):
            continue
        try:
            root = _parse_xml(raw, name)
        except ExtractionError as exc:
            if warnings is not None:
                warnings.append(str(exc))
            continue
        if _local_name(root.tag) not in {"cmLst", "comments"} or not _known_namespace(root.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS})):
            continue
        for comment in root.iter():
            if _local_name(comment.tag) not in {"comment", "cm"}:
                continue
            text_elements = [
                item
                for item in comment.iter()
                if _local_name(item.tag) == "text" and _known_namespace(item.tag)
            ]
            text = "\n".join((item.text or "") for item in text_elements)
            if not text:
                text = _object_text(comment)
            result.append({"part": name, "author_id": _attr(comment, "authorId"), "text": text})
    return result


def _slide_sort_key(name: str) -> tuple[int, str]:
    match = _SLIDE_NUMBER.match(name)
    return (int(match.group(1)), name) if match else (0, name)


def _presentation_part(
    names_to_bytes: dict[str, bytes],
    content_types: dict[str, str],
    by_source: dict[str, list[RelationshipRecord]],
    warnings: list[str],
) -> str | None:
    root_relationship = next(
        (
            relation
            for relation in by_source.get("", [])
            if _relationship_matches(relation, "officeDocument") and relation.resolved_target
        ),
        None,
    )
    if root_relationship is not None:
        if root_relationship.resolved_target in names_to_bytes:
            return root_relationship.resolved_target
        warnings.append(f"Root relationship points to missing presentation part: {root_relationship.resolved_target}")

    if "ppt/presentation.xml" in names_to_bytes:
        if root_relationship is None:
            warnings.append("Using conventional presentation part because the package root relationship is missing")
        return "ppt/presentation.xml"

    presentation_types = (
        "presentationml.presentation.main+xml",
        "presentationml.presentation+xml",
    )
    for name in sorted(names_to_bytes):
        content_type = content_types.get(name, "").casefold()
        if any(content_type.endswith(expected) for expected in presentation_types):
            return name
    return None


def extract_pptx(
    source: str | Path,
    evidence_dir: str | Path | None = None,
    *,
    include_visual_evidence: bool = True,
    include_native_diagrams: bool = True,
) -> ExtractionReport:
    """Extract a PPTX without discarding unsupported OOXML.

    When ``evidence_dir`` is supplied, the original archive is copied as
    ``original.pptx`` and every ZIP member is retained below ``parts/``.
    """
    source_path = Path(source).expanduser().resolve()
    if not source_path.is_file():
        raise ExtractionError(f"PPTX file does not exist: {source_path}")
    try:
        source_bytes = source_path.read_bytes()
    except OSError as exc:
        raise ExtractionError(f"Could not read PPTX file: {source_path}: {exc}") from exc
    warnings: list[str] = []
    try:
        archive = zipfile.ZipFile(BytesIO(source_bytes))
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError) as exc:
        raise ExtractionError(f"Not a readable PPTX ZIP package: {exc}") from exc

    with archive:
        try:
            infos = archive.infolist()
        except Exception as exc:
            raise ExtractionError(f"Could not inspect PPTX ZIP package: {exc}") from exc
        names_to_bytes: dict[str, bytes] = {}
        info_by_name: dict[str, zipfile.ZipInfo] = {}
        for info in infos:
            if info.is_dir():
                _safe_part_name(info.filename, directory=True)
                continue
            name = _safe_part_name(info.filename)
            if name in names_to_bytes:
                raise ExtractionError(f"Duplicate ZIP member: {name}")
            try:
                names_to_bytes[name] = archive.read(info)
                info_by_name[name] = info
            except Exception as exc:
                raise ExtractionError(f"Could not read package part {name}: {exc}") from exc

    if CONTENT_TYPES not in names_to_bytes:
        warnings.append("Missing [Content_Types].xml")
    content_types = _content_types(names_to_bytes.get(CONTENT_TYPES), warnings)
    relationships, by_source = _relationships(names_to_bytes, warnings)
    part_records = [
        PartRecord(
            name=name,
            content_type=_part_content_type(name, content_types),
            size=len(raw),
            compressed_size=info_by_name.get(name).compress_size if name in info_by_name else len(raw),
            crc32=f"{info_by_name.get(name).CRC if name in info_by_name else 0:08x}",
            sha256=_sha256(raw),
            is_xml=name.casefold().endswith(".xml") or name.casefold().endswith(".rels"),
        )
        for name, raw in sorted(names_to_bytes.items())
    ]
    media = [
        MediaRecord(part=item.name, content_type=item.content_type, size=item.size, sha256=item.sha256)
        for item in part_records
        if item.name.startswith("ppt/media/") or item.content_type.casefold().split("/", 1)[0] in {"audio", "image", "video"}
    ]
    asset_targets = {
        relation.resolved_target
        for relation in relationships
        if relation.resolved_target
        and (_relationship_matches(relation, "/image") or _relationship_matches(relation, "/oleObject"))
    }
    embedded_targets = {
        relation.resolved_target
        for relation in relationships
        if relation.resolved_target and _relationship_matches(relation, "/oleObject")
    }
    asset_parts = [
        item
        for item in part_records
        if item.name.startswith("ppt/media/")
        or item.name.startswith("ppt/embeddings/")
        or item.name in asset_targets
    ]
    asset_ids = {item.name: f"asset-{index:04d}" for index, item in enumerate(asset_parts, 1)}

    slides_from_relationships: list[str] = []
    presentation_part = _presentation_part(names_to_bytes, content_types, by_source, warnings)
    if presentation_part is None:
        raise ExtractionError("PPTX package has no presentation part")
    presentation = names_to_bytes.get(presentation_part)
    presentation_root: Any | None = None
    if presentation is not None:
        presentation_root = _parse_xml(presentation, presentation_part)
        if _local_name(presentation_root.tag) != "presentation" or not _known_namespace(presentation_root.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS})):
            raise ExtractionError(f"Invalid presentation part {presentation_part}: expected presentation root")
        presentation_relations = _relationship_map(by_source.get(presentation_part, []))
        for element in presentation_root.iter():
            if _local_name(element.tag) != "sldId" or not _known_namespace(element.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS})):
                continue
            relationship_id = _relationship_attr(element, "id")
            relation = presentation_relations.get(relationship_id or "")
            if relation is None:
                warnings.append(f"Slide entry has no relationship: {relationship_id or '<missing id>'}")
            elif not _relationship_matches(relation, "slide"):
                warnings.append(f"Presentation relationship is not a slide relationship: {relation.relationship_id}")
            elif relation.resolved_target:
                slides_from_relationships.append(relation.resolved_target)
            else:
                warnings.append(f"Slide relationship points to an unresolved part: {relation.relationship_id}")
    discovered_slides = sorted(
        (name for name in names_to_bytes if _SLIDE_NUMBER.match(name)),
        key=_slide_sort_key,
    )
    unique_slides: list[str] = []
    seen_slides: set[str] = set()
    for part in slides_from_relationships:
        if part in seen_slides:
            warnings.append(f"Duplicate slide relationship target: {part}")
            continue
        seen_slides.add(part)
        unique_slides.append(part)
    slides_from_relationships = unique_slides
    if not slides_from_relationships:
        slides_from_relationships = discovered_slides
    elif any(name not in seen_slides for name in discovered_slides):
        warnings.append("Unlisted conventional slide parts were retained as package evidence")

    slides: list[SlideRecord] = []
    canonical_slides: list[dict[str, Any]] = []
    canonical_objects: list[dict[str, Any]] = []
    slide_width, slide_height = _slide_dimensions(presentation_root)
    for part in slides_from_relationships:
        raw = names_to_bytes.get(part)
        if raw is None:
            warnings.append(f"Slide relationship points to missing part: {part}")
            continue
        root = _parse_xml(raw, part)
        if _local_name(root.tag) != "sld" or not _known_namespace(root.tag, frozenset({PRESENTATION_NS, STRICT_PRESENTATION_NS})):
            raise ExtractionError(f"Invalid slide part {part}: expected sld root")
        layout = _rels_to(by_source, part, "/slideLayout")
        master = _rels_to(by_source, layout, "/slideMaster") if layout else None
        theme = _rels_to(by_source, master, "/theme") if master else None
        if layout and layout not in names_to_bytes:
            warnings.append(f"Slide {part} points to missing layout part: {layout}")
        if master and master not in names_to_bytes:
            warnings.append(f"Layout {layout} points to missing master part: {master}")
        if theme and theme not in names_to_bytes:
            warnings.append(f"Master {master} points to missing theme part: {theme}")
        layout_root = _parse_optional_xml(names_to_bytes.get(layout), layout, warnings) if layout else None
        master_root = _parse_optional_xml(names_to_bytes.get(master), master, warnings) if master else None
        theme_root = _parse_optional_xml(names_to_bytes.get(theme), theme, warnings) if theme else None
        notes_part = _rels_to(by_source, part, "/notesSlide")
        if notes_part and notes_part not in names_to_bytes:
            warnings.append(f"Slide {part} points to missing notes part: {notes_part}")
        notes_root = _parse_optional_xml(names_to_bytes.get(notes_part), notes_part, warnings) if notes_part else None
        notes = _texts(notes_root) if notes_root is not None else []
        slide_text = _texts(root)
        slide_hyperlinks = _hyperlinks(root, by_source.get(part, []))
        slide_alt_text = _alt_text(root)
        slide_animations = _animations(root)
        slide_number = len(slides) + 1
        slides.append(
            SlideRecord(
                number=slide_number,
                part=part,
                layout_part=layout,
                master_part=master,
                theme_part=theme,
                text=slide_text,
                notes=notes,
                hyperlinks=slide_hyperlinks,
                alt_text=slide_alt_text,
                animations=slide_animations,
            )
        )
        slide_id = f"slide-{slide_number:02d}"
        slide_objects = _slide_objects(
            root,
            slide_id,
            part,
            by_source.get(part, []),
            slide_width,
            slide_height,
            asset_ids,
            names_to_bytes,
            layout_root,
            master_root,
            theme_root,
        )
        canonical_objects.extend(slide_objects)
        native_ids: set[str] = set()
        for item in slide_objects:
            native_id = item.get("native_id")
            if native_id in (None, ""):
                warnings.append(f"Slide {part} object {item['id']} has no native shape ID")
                continue
            native_id_text = str(native_id)
            if native_id_text in native_ids:
                warnings.append(f"Slide {part} has duplicate native shape ID: {native_id_text}")
            native_ids.add(native_id_text)
            try:
                native_id_number = int(native_id_text, 10)
            except (TypeError, ValueError):
                native_id_number = 0
            if not native_id_text.isascii() or not native_id_text.isdigit() or native_id_number <= 0:
                warnings.append(f"Slide {part} has an invalid native shape ID: {native_id_text!r}")
        canonical_slides.append(
            {
                "id": slide_id,
                "number": slide_number,
                "part": part,
                "layout_part": layout,
                "master_part": master,
                "theme_part": theme,
                "text": slide_text,
                "notes": notes,
                "hyperlinks": slide_hyperlinks,
                "alt_text": slide_alt_text,
                "animations": slide_animations,
                "object_ids": [item["id"] for item in slide_objects],
                "slide_reading_order": "unknown",
                "diagram_flow_direction": "unknown",
                "flow_present": None,
                "flow_presence_basis": "undetermined",
                "visual_region_ids": [],
                "visual_evidence_visibility": {
                    "native": "verified",
                    "rendered": "not_requested",
                    "ocr": "not_requested",
                    "vision": "not_requested",
                },
                "source": _native_source(part, "/sld[1]"),
            }
        )

    canonical_relationships = []
    for relation in relationships:
        relation_id = f"{relation.source_part or 'package'}:{relation.relationship_id}"
        canonical_relationships.append(
            {
                "id": relation_id,
                "source_part": relation.source_part,
                "relationship_id": relation.relationship_id,
                "relationship_type": relation.relationship_type,
                "target": relation.target,
                "target_mode": relation.target_mode,
                "resolved_target": relation.resolved_target,
                "source": _native_source(relation.source_part or None, None),
            }
        )
    canonical_assets = [
        {
            "id": asset_ids[item.name],
            "type": "embedded" if item.name.startswith("ppt/embeddings/") or item.name in embedded_targets else "media",
            "part": item.name,
            "content_type": item.content_type,
            "size": item.size,
            "sha256": item.sha256,
            "source": {
                "layer": "native_ooxml",
                "package_part": item.name,
                "xml_part": None,
                "xml_path": None,
                "confidence": 1.0,
            },
        }
        for item in asset_parts
    ]
    canonical = DeckIR(
        deck={
            "id": "deck",
            "schema": DECKIR_SCHEMA,
            "schema_version": DECKIR_SCHEMA_VERSION,
            "name": source_path.name,
            "source": source_path.name,
            "source_sha256": _sha256(source_bytes),
            "slide_count": len(canonical_slides),
            "slide_size_emu": [slide_width, slide_height],
            "slide_aspect_ratio": round(slide_width / slide_height, 12) if slide_height else None,
            "package_part_count": len(part_records),
        },
        slides=canonical_slides,
        objects=canonical_objects,
        assets=canonical_assets,
        relationships=canonical_relationships,
        rendered_evidence=[],
        ocr_evidence=[],
        vision_evidence=[],
        warnings=warnings,
        provenance={
            "schema": DECKIR_SCHEMA,
            "schema_version": DECKIR_SCHEMA_VERSION,
            "parser_version": VERSION,
            "dependencies": _dependency_versions(),
            "source_name": source_path.name,
            "package_parts": [asdict(item) for item in part_records],
            "authority": {
                "native_ooxml": "authoritative",
                "rendered_cv": "derived",
                "ocr": "derived",
                "vision_model": "probabilistic",
            },
        },
    )
    if include_visual_evidence:
        add_native_visual_evidence(canonical, names_to_bytes)
    if include_native_diagrams:
        add_native_diagram_evidence(canonical)
    canonical.validate()

    report = ExtractionReport(
        source=str(source_path),
        source_sha256=_sha256(source_bytes),
        parser_version=VERSION,
        dependencies=_dependency_versions(),
        package_parts=part_records,
        relationships=relationships,
        slides=slides,
        media=media,
        comments=_comments(names_to_bytes, warnings),
        convenience=_python_pptx_summary(source_path),
        warnings=warnings,
        canonical=canonical,
    )
    if evidence_dir is not None:
        destination = Path(evidence_dir).expanduser().resolve()
        try:
            destination.mkdir(parents=True, exist_ok=True)
            (destination / "parts").mkdir(exist_ok=True)
            (destination / "original.pptx").write_bytes(source_bytes)
            for name, raw in names_to_bytes.items():
                target = destination / "parts" / name
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(raw)
        except OSError as exc:
            raise ExtractionError(f"Could not write evidence bundle: {destination}: {exc}") from exc
        report.evidence_dir = str(destination)
    return report


def extract_document(
    source: str | Path,
    evidence_dir: str | Path | None = None,
    *,
    include_visual_evidence: bool = True,
    include_native_diagrams: bool = True,
    password: str | None = None,
    problem_statement: Any = None,
) -> ExtractionReport:
    """Dispatch a supported source to its native extraction adapter."""
    source_path = Path(source).expanduser().resolve()
    if not source_path.is_file():
        raise ExtractionError(f"Source document does not exist: {source_path}")
    try:
        with source_path.open("rb") as stream:
            signature = stream.read(8)
    except OSError as exc:
        raise ExtractionError(f"Could not read source document: {source_path}: {exc}") from exc
    if signature.startswith(b"%PDF-"):
        from .pdf import extract_pdf

        report = extract_pdf(
            source_path,
            evidence_dir,
            include_visual_evidence=include_visual_evidence,
            include_native_diagrams=include_native_diagrams,
            password=password,
        )
    elif signature.startswith(b"PK"):
        report = extract_pptx(
            source_path,
            evidence_dir,
            include_visual_evidence=include_visual_evidence,
            include_native_diagrams=include_native_diagrams,
        )
    else:
        raise ExtractionError(f"Unsupported source document format: {source_path}")

    if problem_statement is not None and getattr(report, "canonical", None) is not None:
        ps_dict = (
            problem_statement.to_dict()
            if hasattr(problem_statement, "to_dict")
            else dict(problem_statement)
            if isinstance(problem_statement, Mapping)
            else None
        )
        if ps_dict:
            report.canonical.deck["problem_statement"] = ps_dict
    return report
